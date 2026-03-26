from pathlib import Path

from openpyxl import Workbook

from backend.app.parsers.chunking import choose_chunking_strategy
from backend.app.parsers.xlsx_parser import XlsxParser
from backend.app.schemas import BlockRecord, CitationAnchor, NormalizedDocument
from backend.app.parsers.base import segment_to_chunks


def test_xlsx_parser_emits_sheet_region_hints(tmp_path: Path) -> None:
    path = tmp_path / "report.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Revenue"
    sheet.append(["Month", "Value"])
    sheet.append(["Jan", 10])
    sheet.append(["Feb", 12])
    sheet.append([None, None])
    sheet.append(["Owner", "Finance"])
    sheet.append(["Currency", "USD"])
    workbook.save(path)

    document = XlsxParser().parse(path, "report.xlsx")

    hints = document.metadata["chunking_hints"]
    assert hints["preferred_strategy"] == "xlsx_sheet_window"
    assert hints["sheet_count"] == 1
    assert len(hints["sheets"][0]["regions"]) == 2
    assert hints["sheets"][0]["regions"][0]["region_type"] == "table"
    assert hints["sheets"][0]["regions"][1]["region_type"] == "key_value"


def test_choose_chunking_strategy_uses_xlsx_hints() -> None:
    document = NormalizedDocument(
        document_id="doc-xlsx",
        source_path="uploads/report.xlsx",
        file_name="report.xlsx",
        file_type="xlsx",
        title="Workbook",
        checksum="checksum",
        metadata={"chunking_hints": {"preferred_strategy": "xlsx_sheet_window"}},
    )

    strategy = choose_chunking_strategy(document)

    assert strategy.name == "xlsx_sheet_window"


def test_segment_to_chunks_uses_xlsx_structure_aware_strategy() -> None:
    document = NormalizedDocument(
        document_id="doc-xlsx-2",
        source_path="uploads/report.xlsx",
        file_name="report.xlsx",
        file_type="xlsx",
        title="Workbook",
        checksum="checksum",
        parser_used="xlsx_parser",
        parser_confidence=0.85,
        blocks=[
            BlockRecord(
                block_id="h1",
                block_type="heading",
                text="Revenue",
                order_index=0,
                title="Revenue",
                section_path=["Workbook", "Revenue"],
                citation_anchor=CitationAnchor(sheet_name="Revenue", source_label="Revenue"),
                heading_level=1,
            ),
            BlockRecord(
                block_id="h2",
                block_type="heading",
                text="Table 1: Month, Value",
                order_index=1,
                title="Table 1: Month, Value",
                section_path=["Workbook", "Revenue", "Table 1: Month, Value"],
                citation_anchor=CitationAnchor(sheet_name="Revenue", line_start=1, source_label="Revenue:1"),
                heading_level=2,
            ),
            BlockRecord(
                block_id="s1",
                block_type="table",
                text="Headers: Month, Value\n\nSample rows:\nMonth: Jan | Value: 10",
                order_index=2,
                title="Table 1: Month, Value",
                section_path=["Workbook", "Revenue", "Table 1: Month, Value"],
                citation_anchor=CitationAnchor(sheet_name="Revenue", line_start=1, line_end=3, source_label="Revenue:1-3"),
                metadata={"sheet_name": "Revenue", "region_type": "table", "headers": ["Month", "Value"]},
            ),
            BlockRecord(
                block_id="r1",
                block_type="table_row",
                text="Month: Jan | Value: 10",
                order_index=3,
                title="Table 1: Month, Value",
                section_path=["Workbook", "Revenue", "Table 1: Month, Value"],
                citation_anchor=CitationAnchor(sheet_name="Revenue", line_start=2, source_label="Revenue:2"),
                metadata={"sheet_name": "Revenue", "region_type": "table", "headers": ["Month", "Value"], "row_index": 2},
            ),
            BlockRecord(
                block_id="r2",
                block_type="table_row",
                text="Month: Feb | Value: 12",
                order_index=4,
                title="Table 1: Month, Value",
                section_path=["Workbook", "Revenue", "Table 1: Month, Value"],
                citation_anchor=CitationAnchor(sheet_name="Revenue", line_start=3, source_label="Revenue:3"),
                metadata={"sheet_name": "Revenue", "region_type": "table", "headers": ["Month", "Value"], "row_index": 3},
            ),
        ],
        metadata={
            "document_structure": "workbook_sheets",
            "chunking_hints": {
                "preferred_strategy": "xlsx_sheet_window",
                "structure": "workbook_sheets",
                "sheets": [
                    {
                        "sheet_name": "Revenue",
                        "section_path": ["Workbook", "Revenue"],
                        "regions": [
                            {
                                "region_index": 1,
                                "region_type": "table",
                                "region_title": "Table 1: Month, Value",
                                "section_path": ["Workbook", "Revenue", "Table 1: Month, Value"],
                                "summary_block_id": "s1",
                                "row_block_ids": ["r1", "r2"],
                                "headers": ["Month", "Value"],
                                "row_start": 1,
                                "row_end": 3,
                            }
                        ],
                    }
                ],
            },
        },
    )

    chunks, images = segment_to_chunks(document, max_words=80, overlap_words=10)

    assert document.metadata["chunking_strategy"] == "xlsx_sheet_window"
    assert document.metadata["structure_sheet_count"] == 1
    assert len(images) == 0
    assert len(chunks) == 2
    assert chunks[0].chunk_type == "table_chunk"
    assert chunks[1].chunk_type == "row_chunk"
    assert chunks[0].metadata["sheet_name"] == "Revenue"
    assert chunks[1].metadata["row_start"] == 2
    assert "Sheet name: Revenue" in chunks[1].embedding_text
    assert "Headers: Month, Value" in chunks[1].embedding_text
    assert "Month: Feb | Value: 12" in chunks[1].display_text
